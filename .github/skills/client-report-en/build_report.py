"""Build and verify an evidence-based English client report."""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from statistics import fmean, median

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SUPPORTED_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
FONT_NAME = "Arial"
FIGURES = {
    "01_weekly_demand.png": (
        "Weekly order demand",
        "Total meal orders vary substantially across observed weeks. Week numbers are not calendar dates.",
    ),
    "02_top_categories.png": (
        "Most-ordered meal categories",
        "The top three categories represent a large share of total orders, led by Beverages.",
    ),
    "03_center_type_average.png": (
        "Average orders by center type",
        "TYPE_B has the highest average per-center total among the displayed center types; the comparison does not establish why groups differ.",
    ),
    "04_promotion_medians.png": (
        "Observed orders with and without promotion",
        "Promotion exposure is associated with higher observed orders, but matched comparisons are observational and do not prove a causal effect.",
    ),
    "05_discount_order_association.png": (
        "Discount and order association",
        "The measured rank correlation is positive but weak; the association does not demonstrate causation.",
    ),
}


def notebook_text(path: Path) -> tuple[str, str]:
    notebook = json.loads(path.read_text(encoding="utf-8-sig"))
    cells = notebook.get("cells")
    if not isinstance(cells, list):
        raise ValueError(f"Notebook has no valid cells: {path}")
    source_parts: list[str] = []
    output_parts: list[str] = []
    for cell in cells:
        if not isinstance(cell, dict):
            continue
        source = cell.get("source", "")
        source_parts.append("".join(source) if isinstance(source, list) else str(source))
        outputs = cell.get("outputs", [])
        if not isinstance(outputs, list):
            continue
        for output in outputs:
            if not isinstance(output, dict):
                continue
            if output.get("output_type") == "error":
                output_parts.append(
                    f"{output.get('ename', 'Error')}: {output.get('evalue', '')}"
                )
            if output.get("output_type") == "stream":
                text = output.get("text", "")
                output_parts.append("".join(text) if isinstance(text, list) else str(text))
            data = output.get("data", {})
            if isinstance(data, dict) and "text/plain" in data:
                text = data["text/plain"]
                output_parts.append("".join(text) if isinstance(text, list) else str(text))
    return "\n".join(source_parts), "\n".join(output_parts)


def required(pattern: str, text: str, label: str, flags: int = re.IGNORECASE) -> re.Match[str]:
    match = re.search(pattern, text, flags)
    if match is None:
        raise ValueError(
            f"Could not extract {label} from saved notebook outputs. "
            "Run the relevant notebook cells and save their outputs."
        )
    return match


def extract_results(analysis_source: str, analysis_output: str, regression_output: str) -> dict[str, str]:
    analysis = analysis_source + "\n" + analysis_output
    weekly = required(r"Average weekly orders:\s*([\d,]+)", analysis, "average weekly orders")
    high = required(
        r"Highest-demand week:\s*(\d+)\s*\(([\d,]+)\s*orders\)",
        analysis,
        "highest-demand week",
    )
    low = required(
        r"Lowest-demand week:\s*(\d+)\s*\(([\d,]+)\s*orders\)",
        analysis,
        "lowest-demand week",
    )
    top_three = required(
        r"three largest categories account for\s*([\d.]+)%",
        analysis,
        "share of the top three categories",
    )
    beverages = required(r"Beverages\s+([\d,]+)\s+([\d.]+)", analysis_output, "Beverages category")
    rho = required(r"spearman_rho\s+([+-]?[\d.]+)", analysis_output, "discount rank correlation")
    observations = required(
        r"valid_observations\s+([\d,]+)", analysis_output, "discount analysis observations"
    )
    rmse = required(
        r"Root mean squared error\s*\(orders\):\s*([\d,.]+)",
        regression_output,
        "validation RMSE",
    )
    mae = required(
        r"Mean absolute error\s*\(orders\):\s*([\d,.]+)",
        regression_output,
        "validation MAE",
    )
    validation = required(
        r"Validation weeks:\s*(\d+)\s*to\s*(\d+)", regression_output, "validation weeks"
    )
    validation_rows = required(
        r"Validation rows:\s*([\d,]+)", regression_output, "validation rows"
    )
    train = required(
        r"Training rows and columns:\s*\(([\d,]+),\s*(\d+)\)",
        regression_output,
        "training row count",
    )
    test = required(
        r"Test rows and columns:\s*\(([\d,]+),\s*(\d+)\)",
        regression_output,
        "test row count",
    )
    train_weeks = required(
        r"Training weeks:\s*(\d+)\s*to\s*(\d+)", regression_output, "training week range"
    )
    test_weeks = required(
        r"Test weeks:\s*(\d+)\s*to\s*(\d+)", regression_output, "test week range"
    )

    summary = required(
        r"Center 13 leads total orders \(([\d,]+)\).*?"
        r"TYPE_B has the highest average total per center "
        r"\(([\d,]+) across (\d+) centers\)",
        analysis_source,
        "center and center-type summary",
        re.IGNORECASE | re.DOTALL,
    )
    cuisines = required(r"Italian leads \(([\d.]+)%\)", analysis_source, "leading cuisine share")
    promotions = required(
        r"unadjusted means are higher for promoted observations:\s*"
        r"email ([\d.]+) versus ([\d.]+) without email;\s*"
        r"homepage-featured ([\d.]+) versus ([\d.]+) without featuring",
        analysis_source,
        "promotion comparison",
        re.IGNORECASE,
    )
    return {
        "average_weekly": weekly.group(1),
        "high_week": high.group(1),
        "high_orders": high.group(2),
        "low_week": low.group(1),
        "low_orders": low.group(2),
        "top_three_share": top_three.group(1),
        "beverage_orders": beverages.group(1),
        "beverage_share": beverages.group(2),
        "rho": rho.group(1),
        "discount_observations": observations.group(1),
        "rmse": rmse.group(1),
        "mae": mae.group(1),
        "validation_weeks": f"{validation.group(1)}–{validation.group(2)}",
        "validation_rows": f"{int(validation_rows.group(1).replace(',', '')):,}",
        "train_rows": f"{int(train.group(1).replace(',', '')):,}",
        "test_rows": f"{int(test.group(1).replace(',', '')):,}",
        "train_weeks": f"{train_weeks.group(1)}–{train_weeks.group(2)}",
        "test_weeks": f"{test_weeks.group(1)}–{test_weeks.group(2)}",
        "top_center_orders": summary.group(1),
        "type_b_center_total": summary.group(2),
        "type_b_centers": summary.group(3),
        "italian_share": cuisines.group(1),
        "email_with": promotions.group(1),
        "email_without": promotions.group(2),
        "homepage_with": promotions.group(3),
        "homepage_without": promotions.group(4),
    }


def target_statistics(path: Path) -> dict[str, float]:
    values: list[float] = []
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.DictReader(csv_file)
        if not reader.fieldnames or "num_orders" not in reader.fieldnames:
            raise ValueError(f"Required num_orders column is missing from {path}")
        for row in reader:
            value = row.get("num_orders")
            if value not in (None, ""):
                values.append(float(value))
    if not values:
        raise ValueError(f"No valid num_orders values found in {path}")
    return {
        "count": float(len(values)),
        "minimum": min(values),
        "maximum": max(values),
        "mean": fmean(values),
        "median": median(values),
    }


def set_ltr_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)
    bidi.set(qn("w:val"), "0")


def set_run(run, size: int = 12, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def add_paragraph(document: Document, text: str, size: int = 12, bold: bool = False):
    paragraph = document.add_paragraph()
    set_ltr_paragraph(paragraph)
    set_run(paragraph.add_run(text), size, bold)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    set_ltr_paragraph(paragraph)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    set_run(paragraph.add_run(text), 20 if level == 1 else 18, True)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    set_ltr_paragraph(paragraph)
    set_run(paragraph.add_run(text), 12)


def add_figure(document: Document, path: Path, width: float = 6.2) -> None:
    title, caption = FIGURES.get(
        path.name,
        ("Additional analysis figure", f"Additional analysis result shown in {path.stem}."),
    )
    paragraph = document.add_paragraph()
    set_ltr_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    set_ltr_paragraph(caption_paragraph)
    set_run(caption_paragraph.add_run(f"Figure: {title}. {caption}"), 10)


def add_model_table(document: Document, results: dict[str, str]) -> None:
    rows = [
        ("Measure", "Result"),
        ("Training records", f"{results['train_rows']} (weeks {results['train_weeks']})"),
        ("Test records", f"{results['test_rows']} (weeks {results['test_weeks']})"),
        ("Validation records", f"{results['validation_rows']} (weeks {results['validation_weeks']})"),
        ("Mean absolute error (MAE)", f"{results['mae']} orders per center-meal-week record"),
        ("Root mean squared error (RMSE)", f"{results['rmse']} orders per center-meal-week record"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    visual = table._tbl.tblPr.find(qn("w:bidiVisual"))
    if visual is None:
        visual = OxmlElement("w:bidiVisual")
        table._tbl.tblPr.append(visual)
    visual.set(qn("w:val"), "0")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                set_ltr_paragraph(paragraph)
                for run in paragraph.runs:
                    set_run(run, 11, row_index == 0)


def add_page_number(paragraph) -> None:
    set_ltr_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    run = paragraph.add_run("Page ")
    set_run(run, 10)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    set_ltr_paragraph(paragraph)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def make_report(project_root: Path) -> tuple[Path, str, int]:
    analysis_path = project_root / "notebooks" / "analysis.ipynb"
    regression_path = project_root / "notebooks" / "linear_regression.ipynb"
    if not regression_path.is_file():
        regression_path = project_root / "notebooks" / "lenier_regression.ipynb"
    figures_directory = project_root / "reports" / "figures"
    training_csv = project_root / "data" / "raw" / "train.csv"
    for required_path in (analysis_path, regression_path, figures_directory, training_csv):
        if not required_path.exists():
            raise FileNotFoundError(f"Required input was not found: {required_path}")

    figures = sorted(path for path in figures_directory.iterdir() if path.is_file())
    unsupported = [path.name for path in figures if path.suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES]
    if unsupported:
        raise ValueError(f"Unsupported figure format(s); none were skipped: {', '.join(unsupported)}")
    if not figures:
        raise FileNotFoundError(f"No figures found in {figures_directory}")

    analysis_source, analysis_output = notebook_text(analysis_path)
    _, regression_output = notebook_text(regression_path)
    results = extract_results(analysis_source, analysis_output, regression_output)
    stats = target_statistics(training_csv)
    rmse_number = float(results["rmse"].replace(",", ""))
    mean_orders = stats["mean"]
    relative_rmse = rmse_number / mean_orders * 100

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(12)
    normal_properties = normal._element.get_or_add_rPr()
    normal_fonts = normal_properties.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_properties.insert(0, normal_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    cover = document.add_paragraph()
    set_ltr_paragraph(cover, WD_ALIGN_PARAGRAPH.CENTER)
    cover.paragraph_format.space_before = Pt(120)
    set_run(cover.add_run("Food Demand Analysis"), 24, True)
    subtitle = document.add_paragraph()
    set_ltr_paragraph(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    set_run(subtitle.add_run("Client Report | Findings and Recommendations"), 18, True)
    add_paragraph(document, "Prepared for business decision-makers", 14, True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_break(document)

    add_heading(document, "Executive Summary")
    add_paragraph(
        document,
        f"Across 145 observed weeks, average weekly demand was {results['average_weekly']} orders. "
        f"The highest weekly total was {results['high_orders']} in week {results['high_week']}; "
        f"the lowest was {results['low_orders']} in week {results['low_week']}. "
        f"The three largest meal categories represented {results['top_three_share']}% of orders. "
        f"A linear-regression validation on weeks {results['validation_weeks']} produced an RMSE of "
        f"{results['rmse']} orders per center-meal-week record. The evidence can inform demand review, "
        "but the observed promotion and price relationships are not causal findings.",
    )

    add_heading(document, "Objective and Data")
    add_paragraph(
        document,
        "This analysis describes weekly demand, meal and center patterns, observed relationships with "
        "promotion and discount, and the validation performance of a linear-regression model. The "
        f"prepared analysis dataset contains 456,548 rows, 145 weeks, 77 centers, and 51 meals, with "
        "no missing cells or duplicate rows reported in the saved notebook output. The regression "
        f"used {results['train_rows']} training records from weeks {results['train_weeks']} and "
        f"predicted {results['test_rows']} unlabeled test records for weeks {results['test_weeks']}.",
    )
    add_paragraph(
        document,
        f"In the read-only training CSV, num_orders ranges from {stats['minimum']:,.0f} to "
        f"{stats['maximum']:,.0f} orders per center-meal-week record; the mean is "
        f"{mean_orders:,.2f} and the median is {stats['median']:,.0f}. This record-level target "
        "must not be confused with the total number of orders across a company in a week.",
    )

    add_heading(document, "Key Findings")
    add_heading(document, "1. Weekly demand varies substantially", 2)
    add_paragraph(
        document,
        f"Average demand was {results['average_weekly']} orders per week. The observed weekly total "
        f"ranged from {results['low_orders']} in week {results['low_week']} to "
        f"{results['high_orders']} in week {results['high_week']}. These are weekly aggregates, "
        "not individual meal-record counts. Week numbers have no calendar dates, so the pattern "
        "cannot be assigned to a season or holiday from these data.",
    )
    if any(path.name == "01_weekly_demand.png" for path in figures):
        add_figure(document, next(path for path in figures if path.name == "01_weekly_demand.png"))

    add_heading(document, "2. A small number of categories account for much of demand", 2)
    add_paragraph(
        document,
        f"The three largest categories account for {results['top_three_share']}% of orders. "
        f"Beverages lead, with {results['beverage_orders']} orders "
        f"({results['beverage_share']}%); the largest individual meal in the notebook summary is "
        "meal 2290, an Indian Rice Bowl, with 9,770,303 orders (8.17%). Italian cuisine has the "
        f"largest cuisine share at {results['italian_share']}%. These are cumulative historical "
        "counts, not a guarantee of future demand.",
    )
    if any(path.name == "02_top_categories.png" for path in figures):
        add_figure(document, next(path for path in figures if path.name == "02_top_categories.png"))

    add_heading(document, "3. Center totals and center-type averages differ", 2)
    add_paragraph(
        document,
        f"Center 13 recorded the highest cumulative total, {results['top_center_orders']} orders. "
        f"TYPE_B had the highest average total per center: {results['type_b_center_total']} across "
        f"{results['type_b_centers']} centers. The total by center type is affected by the different "
        "number of centers in each group; observed differences do not identify their cause.",
    )
    if any(path.name == "03_center_type_average.png" for path in figures):
        add_figure(document, next(path for path in figures if path.name == "03_center_type_average.png"))

    add_heading(document, "4. Higher observed orders coincide with promotion exposure", 2)
    add_paragraph(
        document,
        f"Mean orders were {results['email_with']} for email-promoted records versus "
        f"{results['email_without']} without email promotion; homepage-featured records averaged "
        f"{results['homepage_with']} versus {results['homepage_without']} without featuring. "
        "These are unadjusted associations: promotions may be assigned to different meals or "
        "circumstances, so the comparisons do not show that promotion caused the difference. "
        "The matched center-meal statistical comparisons are exploratory because weeks repeat "
        "within pairs.",
    )
    if any(path.name == "04_promotion_medians.png" for path in figures):
        add_figure(document, next(path for path in figures if path.name == "04_promotion_medians.png"))

    add_heading(document, "5. Discount has a weak positive rank association with orders", 2)
    add_paragraph(
        document,
        f"The Spearman rank correlation between discount percentage and orders is "
        f"{float(results['rho']):.3f}, calculated over {results['discount_observations']} records. "
        "This is a weak positive association, not evidence that larger discounts cause more orders. "
        "Differences in meal, center, and week may also relate to the observed counts.",
    )
    if any(path.name == "05_discount_order_association.png" for path in figures):
        add_figure(document, next(path for path in figures if path.name == "05_discount_order_association.png"))

    known = set(FIGURES)
    extras = [path for path in figures if path.name not in known]
    if extras:
        add_heading(document, "Additional Results")
        for figure in extras:
            add_figure(document, figure)

    add_heading(document, "Predictive Model Results")
    add_paragraph(
        document,
        f"The linear-regression model was evaluated chronologically on the final 10 labeled weeks, "
        f"weeks {results['validation_weeks']}, covering {results['validation_rows']} records. "
        f"Validation RMSE was {results['rmse']} orders per center-meal-week record, and MAE was "
        f"{results['mae']} orders. The model was then refit on all labeled weeks before generating "
        "predictions for test weeks 146–155; actual labels for those test rows are not available "
        "for calculating their prediction error.",
    )
    add_model_table(document, results)

    add_heading(document, "How to Interpret RMSE", 2)
    add_paragraph(
        document,
        f"RMSE expresses prediction error in the target's units: orders per center-meal-week record. "
        f"The RMSE of {results['rmse']} is about {relative_rmse:.1f}% of the training-data mean of "
        f"{mean_orders:,.2f} orders per record. RMSE squares errors before averaging and taking the "
        "square root, so large errors receive more weight; it is not a simple arithmetic average "
        "of the errors. The training target ranges from "
        f"{stats['minimum']:,.0f} to {stats['maximum']:,.0f} orders per record. There is no universal "
        "commercial acceptance threshold: the business must define a tolerable error for each "
        "operational decision and compare this model with a suitable baseline.",
    )
    add_paragraph(
        document,
        f"Hypothetical illustration, not an observed prediction: if one center-meal-week record "
        f"had an actual count of 500 orders, an error with a magnitude of roughly the RMSE "
        f"({rmse_number:.0f} orders) would correspond to a prediction around "
        f"{500-rmse_number:.0f} or {500+rmse_number:.0f} orders. This example explains the scale "
        "of the metric; it does not describe a specific model prediction. It is a per-record "
        "illustration, not an estimate of total weekly company orders.",
    )

    add_heading(document, "Practical Recommendations")
    add_bullet(document, "Pilot forecasts at the center-meal level and compare each forecast with actual orders before using it to adjust supply.")
    add_bullet(document, "Prioritize availability reviews for the largest-demand categories and meals, while checking inventory and waste records separately.")
    add_bullet(document, "Define acceptable forecast error by operational use, then compare the model with a simple baseline on the same chronological validation window.")
    add_bullet(document, "Evaluate promotions and price changes with a controlled experiment or stronger causal design before treating observed associations as expected lift.")

    add_heading(document, "Notes and Limitations")
    add_bullet(document, "The order data do not directly measure inventory, food waste, or stockouts; separate operational records are required.")
    add_bullet(document, "Week numbers are not calendar dates, so these results cannot identify seasonality, holidays, or specific calendar events.")
    add_bullet(document, "Promotion and discount results are observational associations and do not prove causal effects.")
    add_bullet(document, "RMSE and MAE are from labeled training-period validation, not from future test weeks 146–155.")
    add_bullet(document, "The report summarizes aggregated demand patterns and does not attempt to identify the company, customers, or individuals.")

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            set_ltr_paragraph(paragraph)
        add_page_number(section.footer.paragraphs[0])

    output_path = project_root / "reports" / "client_report_en.docx"
    document.save(output_path)
    verify_report(output_path, len(figures), results["rmse"])
    return output_path, results["rmse"], len(figures)


def verify_report(path: Path, expected_images: int, rmse_text: str) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"Report was not created: {path}")
    document = Document(path)
    if len(document.inline_shapes) != expected_images:
        raise RuntimeError(
            f"Embedded image count ({len(document.inline_shapes)}) does not match "
            f"figure-file count ({expected_images})."
        )
    paragraphs = list(document.paragraphs)
    full_text = "\n".join(paragraph.text for paragraph in paragraphs)
    for table in document.tables:
        full_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    if rmse_text not in full_text or "Hypothetical illustration" not in full_text:
        raise RuntimeError("RMSE value or its hypothetical-example explanation is missing.")
    if "Hypothetical illustration" not in full_text or "not an observed prediction" not in full_text:
        raise RuntimeError("The numerical example is not clearly labeled hypothetical.")
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        if paragraph.alignment != WD_ALIGN_PARAGRAPH.LEFT and paragraph.text.strip() not in {
            "Food Demand Analysis",
            "Client Report | Findings and Recommendations",
            "Prepared for business decision-makers",
        }:
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                raise RuntimeError("A report paragraph is not left-aligned or intentionally centered.")
        bidi = paragraph._p.get_or_add_pPr().find(qn("w:bidi"))
        if bidi is None or bidi.get(qn("w:val")) != "0":
            raise RuntimeError("A report paragraph does not have left-to-right text direction.")
    for table in document.tables:
        visual = table._tbl.tblPr.find(qn("w:bidiVisual"))
        if visual is None or visual.get(qn("w:val")) != "0":
            raise RuntimeError("A report table does not have left-to-right table direction.")
        if table.alignment != WD_TABLE_ALIGNMENT.LEFT:
            raise RuntimeError("A report table is not left-aligned.")
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                        raise RuntimeError("A table-cell paragraph is not left-aligned.")
                    bidi = paragraph._p.get_or_add_pPr().find(qn("w:bidi"))
                    if bidi is None or bidi.get(qn("w:val")) != "0":
                        raise RuntimeError("A table-cell paragraph is not left-to-right.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Create and verify the English client report.")
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[3],
        help="Project root directory (defaults to the repository root).",
    )
    args = parser.parse_args()
    output, rmse, image_count = make_report(args.project_root.resolve())
    print(f"Created and verified report: {output}")
    print(f"Validation RMSE: {rmse} orders per center-meal-week record")
    print(f"Embedded figures: {image_count}")


if __name__ == "__main__":
    main()
