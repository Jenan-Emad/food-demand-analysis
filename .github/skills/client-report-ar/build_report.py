"""Build an evidence-based Arabic client report from the project notebooks."""

from __future__ import annotations

import argparse
import csv
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from statistics import fmean, median

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


IMAGE_CAPTIONS = {
    "01_weekly_demand.png": (
        "تغير إجمالي الطلب أسبوعيًا",
        "يوضح الرسم تفاوت إجمالي الطلب بين الأسابيع المرصودة، مع إبراز أعلى أسبوع وأدناه. "
        "أرقام الأسابيع ليست تواريخ تقويمية.",
    ),
    "02_top_categories.png": (
        "أكثر فئات الوجبات طلبًا",
        "تقارن الأعمدة إجمالي الطلبات في أكبر فئات الوجبات، وتظهر تصدّر المشروبات ضمن الفئات المعروضة.",
    ),
    "03_center_type_average.png": (
        "متوسط إجمالي الطلب حسب نوع المركز",
        "يعرض الرسم متوسط الطلب التراكمي لكل مركز حسب نوعه؛ وهو وصف للبيانات وليس تفسيرًا لسبب الاختلاف.",
    ),
    "04_promotion_medians.png": (
        "الطلبات الوسيطة مع الترويج ومن دونه",
        "يعرض مقارنة وصفية داخل أزواج المركز والوجبة. لا تثبت الفروق أن الترويج هو سبب زيادة الطلب.",
    ),
    "05_discount_order_association.png": (
        "العلاقة بين نسبة الخصم والطلبات",
        "يعرض متوسط الطلب عبر شرائح الخصم ومعامل ارتباط الرتب. الارتباط لا يثبت السببية.",
    ),
}

SUPPORTED_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
FONT_NAME = "Arial"


class TextExtractor(HTMLParser):
    """Collect readable text from HTML table outputs in a notebook."""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)


def readable_value(value: object) -> str:
    if isinstance(value, list):
        return "".join(str(part) for part in value)
    return str(value)


def notebook_cell_text(cell: dict[str, object]) -> str:
    parts = [readable_value(cell.get("source", ""))]
    outputs = cell.get("outputs", [])
    if not isinstance(outputs, list):
        return "\n".join(parts)

    for output in outputs:
        if not isinstance(output, dict):
            continue
        if output.get("output_type") == "stream":
            parts.append(readable_value(output.get("text", "")))
        elif output.get("output_type") == "error":
            parts.append(f"{output.get('ename', 'Error')}: {output.get('evalue', '')}")

        data = output.get("data", {})
        if isinstance(data, dict):
            if "text/plain" in data:
                parts.append(readable_value(data["text/plain"]))
            if "text/html" in data:
                parser = TextExtractor()
                parser.feed(readable_value(data["text/html"]))
                parts.append(" ".join(parser.parts))
    return "\n".join(parts)


def read_notebook(path: Path) -> tuple[dict[str, object], dict[str, str]]:
    notebook = json.loads(path.read_text(encoding="utf-8-sig"))
    cells = notebook.get("cells")
    if not isinstance(cells, list):
        raise ValueError(f"دفتر غير صالح أو بلا خلايا: {path}")

    cell_text: dict[str, str] = {}
    for index, cell in enumerate(cells):
        if not isinstance(cell, dict):
            continue
        cell_text[str(cell.get("id", index))] = notebook_cell_text(cell)
    return notebook, cell_text


def required_match(pattern: str, text: str, description: str) -> re.Match[str]:
    match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    if match is None:
        raise ValueError(f"تعذر استخراج {description} من نتائج الدفتر. شغّل خلاياه واحفظ المخرجات.")
    return match


def extract_results(
    analysis_cells: dict[str, str], regression_cells: dict[str, str]
) -> dict[str, object]:
    analysis_text = "\n".join(analysis_cells.values())
    regression_text = "\n".join(regression_cells.values())

    average_weekly = required_match(
        r"Average weekly orders:\s*([\d,]+)", analysis_text, "متوسط الطلب الأسبوعي"
    ).group(1)
    high = required_match(
        r"Highest-demand week:\s*(\d+)\s*\(([\d,]+)\s*orders\)",
        analysis_text,
        "أعلى أسبوع طلبًا",
    )
    low = required_match(
        r"Lowest-demand week:\s*(\d+)\s*\(([\d,]+)\s*orders\)",
        analysis_text,
        "أدنى أسبوع طلبًا",
    )
    top_three_share = required_match(
        r"three largest categories account for\s*([\d.]+)%",
        analysis_text,
        "حصة أكبر ثلاث فئات",
    ).group(1)
    beverage = required_match(
        r"Beverages\s+([\d,]+)\s+([\d.]+)",
        analysis_text,
        "نتيجة فئة المشروبات",
    )
    rho = required_match(
        r"spearman_rho\s+([+-]?[\d.]+)", analysis_text, "ارتباط الخصم والطلبات"
    ).group(1)
    valid_discount_rows = required_match(
        r"valid_observations\s+([\d,]+)", analysis_text, "عدد مشاهدات تحليل الخصم"
    ).group(1)
    rmse = required_match(
        r"Root mean squared error\s*\(orders\):\s*([\d,.]+)",
        regression_text,
        "قيمة RMSE",
    ).group(1)

    mae_match = re.search(
        r"Mean absolute error\s*\(orders\):\s*([\d,.]+)",
        regression_text,
        flags=re.IGNORECASE,
    )
    validation_match = re.search(
        r"Validation weeks:\s*(\d+)\s*to\s*(\d+)", regression_text, flags=re.IGNORECASE
    )
    validation_rows_match = re.search(
        r"Validation rows:\s*([\d,]+)", regression_text, flags=re.IGNORECASE
    )
    train_match = re.search(
        r"Training rows and columns:\s*\(([\d,]+),\s*(\d+)\)",
        regression_text,
        flags=re.IGNORECASE,
    )
    test_match = re.search(
        r"Test rows and columns:\s*\(([\d,]+),\s*(\d+)\)",
        regression_text,
        flags=re.IGNORECASE,
    )
    train_weeks = re.search(
        r"Training weeks:\s*(\d+)\s*to\s*(\d+)", regression_text, flags=re.IGNORECASE
    )
    test_weeks = re.search(
        r"Test weeks:\s*(\d+)\s*to\s*(\d+)", regression_text, flags=re.IGNORECASE
    )
    if not all((train_match, test_match, train_weeks, test_weeks)):
        raise ValueError("مخرجات الانحدار الخطي لا تتضمن صفوف التدريب والاختبار ومدى الأسابيع.")

    return {
        "average_weekly": average_weekly,
        "high_week": high.group(1),
        "high_orders": high.group(2),
        "low_week": low.group(1),
        "low_orders": low.group(2),
        "top_three_share": top_three_share,
        "beverage_orders": beverage.group(1),
        "beverage_share": beverage.group(2),
        "rho": rho,
        "valid_discount_rows": valid_discount_rows,
        "rmse": rmse,
        "mae": mae_match.group(1) if mae_match else None,
        "validation_weeks": (
            f"{validation_match.group(1)}–{validation_match.group(2)}"
            if validation_match
            else "غير متاح"
        ),
        "validation_rows": validation_rows_match.group(1) if validation_rows_match else "غير متاح",
        "train_rows": train_match.group(1),
        "train_week_range": f"{train_weeks.group(1)}–{train_weeks.group(2)}",
        "test_rows": test_match.group(1),
        "test_week_range": f"{test_weeks.group(1)}–{test_weeks.group(2)}",
    }


def read_target_statistics(path: Path) -> dict[str, float]:
    orders: list[float] = []
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.DictReader(csv_file)
        if not reader.fieldnames or "num_orders" not in reader.fieldnames:
            raise ValueError(f"العمود num_orders غير موجود في {path}")
        for row in reader:
            raw_value = row.get("num_orders")
            if raw_value in (None, ""):
                continue
            orders.append(float(raw_value))
    if not orders:
        raise ValueError(f"لا توجد قيم num_orders صالحة في {path}")
    return {
        "count": float(len(orders)),
        "minimum": min(orders),
        "maximum": max(orders),
        "mean": fmean(orders),
        "median": median(orders),
    }


def set_run_font(run, size: int, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    complex_script = OxmlElement("w:cs")
    complex_script.set(qn("w:val"), "1")
    r_pr.append(complex_script)


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_table_rtl(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table_properties = table._tbl.tblPr
    bidi_visual = table_properties.find(qn("w:bidiVisual"))
    if bidi_visual is None:
        bidi_visual = OxmlElement("w:bidiVisual")
        table_properties.append(bidi_visual)
    bidi_visual.set(qn("w:val"), "1")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_paragraph_rtl(paragraph)


def set_document_rtl(document: Document) -> None:
    for paragraph in document.paragraphs:
        set_paragraph_rtl(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)
    for table in document.tables:
        set_table_rtl(table)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            set_paragraph_rtl(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)
        for paragraph in section.footer.paragraphs:
            set_paragraph_rtl(paragraph, paragraph.alignment or WD_ALIGN_PARAGRAPH.RIGHT)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("الصفحة ")
    set_run_font(run, 10)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def add_text(document: Document, text: str, size: int = 12, bold: bool = False):
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    run = paragraph.add_run(text)
    set_run_font(run, size=20 if level == 1 else 16, bold=True)
    return paragraph


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    set_paragraph_rtl(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, 12)


def add_figure(document: Document, image_path: Path, max_width: float = 6.0) -> None:
    title, caption = IMAGE_CAPTIONS.get(
        image_path.name,
        ("رسم من نتائج التحليل", f"يعرض هذا الرسم نتيجة تحليل مرتبطة بالبيانات: {image_path.stem}."),
    )
    paragraph = document.add_paragraph()
    set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    paragraph.add_run().add_picture(str(image_path), width=Inches(max_width))

    caption_paragraph = document.add_paragraph()
    set_paragraph_rtl(caption_paragraph)
    caption_run = caption_paragraph.add_run(f"الشكل: {title}. {caption}")
    set_run_font(caption_run, 11)


def add_summary_table(document: Document, results: dict[str, object]) -> None:
    rows = [
        ("بيانات التدريب", f"{results['train_rows']} صفًا؛ الأسابيع {results['train_week_range']}"),
        ("بيانات الاختبار", f"{results['test_rows']} صفًا؛ الأسابيع {results['test_week_range']}"),
        ("صفوف التحقق", f"{results['validation_rows']}؛ الأسابيع {results['validation_weeks']}"),
        ("متوسط الخطأ المطلق (MAE)", f"{results['mae']} طلبًا" if results["mae"] else "غير متاح"),
        ("جذر متوسط مربع الخطأ (RMSE)", f"{results['rmse']} طلبًا"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "المقياس"
    table.rows[0].cells[1].text = "النتيجة"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    set_table_rtl(table)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, 11, bold=(row_index == 0))


def make_report(project_root: Path) -> Path:
    analysis_path = project_root / "notebooks" / "analysis.ipynb"
    regression_candidates = [
        project_root / "notebooks" / "linear_regression.ipynb",
        project_root / "notebooks" / "lenier_regression.ipynb",
    ]
    regression_path = next((path for path in regression_candidates if path.is_file()), None)
    figures_directory = project_root / "reports" / "figures"
    training_csv = project_root / "data" / "raw" / "train.csv"
    for required_path in (analysis_path, figures_directory, training_csv):
        if not required_path.exists():
            raise FileNotFoundError(f"المسار المطلوب غير موجود: {required_path}")
    if regression_path is None:
        raise FileNotFoundError("لم يُعثر على دفتر نتائج الانحدار الخطي في notebooks/.")

    figure_paths = sorted(
        path for path in figures_directory.iterdir() if path.is_file()
    )
    if not figure_paths:
        raise FileNotFoundError(f"لم توجد صور في {figures_directory}")
    unsupported = [path.name for path in figure_paths if path.suffix.lower() not in SUPPORTED_IMAGE_SUFFIXES]
    if unsupported:
        raise ValueError(f"صيغة صورة غير مدعومة؛ لم يتم تجاوزها: {', '.join(unsupported)}")

    _, analysis_cells = read_notebook(analysis_path)
    _, regression_cells = read_notebook(regression_path)
    results = extract_results(analysis_cells, regression_cells)
    target_stats = read_target_statistics(training_csv)
    rmse_value = float(str(results["rmse"]).replace(",", ""))
    error_ratio = rmse_value / target_stats["mean"] * 100
    mean_target = target_stats["mean"]

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal_style = document.styles["Normal"]
    normal_style.font.name = FONT_NAME
    normal_style.font.size = Pt(12)
    normal_r_pr = normal_style._element.get_or_add_rPr()
    normal_fonts = normal_r_pr.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_r_pr.insert(0, normal_fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    cover_title = document.add_paragraph()
    set_paragraph_rtl(cover_title, WD_ALIGN_PARAGRAPH.CENTER)
    cover_title.paragraph_format.space_before = Pt(120)
    title_run = cover_title.add_run("تقرير تحليل الطلب والتنبؤ بالطلبات")
    set_run_font(title_run, 24, bold=True)
    subtitle = document.add_paragraph()
    set_paragraph_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    subtitle_run = subtitle.add_run("ملخص النتائج والتوصيات")
    set_run_font(subtitle_run, 18, bold=True)
    add_text(document, "تقرير موجه للعميل", size=14, bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    add_heading(document, "الملخص التنفيذي")
    add_text(
        document,
        f"تتفاوت الطلبات الأسبوعية بوضوح؛ بلغ متوسطها {results['average_weekly']} طلبًا، "
        f"وسُجل أعلى إجمالي في الأسبوع {results['high_week']} ({results['high_orders']} طلبًا) "
        f"وأدنى إجمالي في الأسبوع {results['low_week']} ({results['low_orders']} طلبًا). "
        f"مثّلت أكبر ثلاث فئات {results['top_three_share']}% من الطلبات. "
        f"في اختبار زمني على الأسابيع {results['validation_weeks']}، بلغ RMSE "
        f"{results['rmse']} طلبًا لكل سجل مركز-وجبة. هذه مؤشرات للتخطيط والمراجعة، "
        "وليست دليلًا على أن السعر أو الترويج سبب تغير الطلب.",
    )

    add_heading(document, "الهدف والبيانات")
    add_text(
        document,
        f"يهدف التحليل إلى وصف تغير الطلب ومزيج الوجبات والعلاقة المرصودة مع الترويج والخصم، "
        f"ثم اختبار نموذج خطي للتنبؤ بعدد الطلبات. شمل التدريب {results['train_rows']} صفًا "
        f"للأسابيع {results['train_week_range']}، ويشمل ملف الاختبار {results['test_rows']} صفًا "
        f"للأسابيع {results['test_week_range']}. كل صف في بيانات التدريب والاختبار يصف وجبة "
        "في مركز خلال أسبوع؛ لذا فإن خطأ النموذج ليس خطأ مجموع طلبات الشركة الأسبوعي.",
    )
    add_text(
        document,
        "مقياس الطلب في بيانات التدريب تراوح بين "
        f"{target_stats['minimum']:,.0f} و{target_stats['maximum']:,.0f} طلبًا لكل سجل، "
        f"ومتوسطه {mean_target:,.2f} ووسيطه {target_stats['median']:,.0f}.",
    )

    add_heading(document, "أهم النتائج")
    add_heading(document, "1. تذبذب الطلب الأسبوعي", level=2)
    add_text(
        document,
        f"تراوحت المجاميع المرصودة بين {results['low_orders']} طلبًا في الأسبوع "
        f"{results['low_week']} و{results['high_orders']} طلبًا في الأسبوع {results['high_week']}. "
        "يساعد ذلك على توجيه الانتباه إلى الأسابيع ذات الطلب غير المعتاد، لكن غياب تواريخ التقويم "
        "لا يسمح بربط القمم بموسم أو مناسبة.",
    )
    if (figure := next((path for path in figure_paths if path.name == "01_weekly_demand.png"), None)):
        add_figure(document, figure)

    add_heading(document, "2. تركّز الطلب في فئات محددة", level=2)
    add_text(
        document,
        f"استحوذت أكبر ثلاث فئات على {results['top_three_share']}% من الطلبات. "
        f"وسجلت فئة المشروبات {results['beverage_orders']} طلبًا "
        f"({results['beverage_share']}% من الإجمالي)، ما يجعل متابعة هذه الفئات مفيدة عند "
        "ترتيب أولويات مراجعة التوقعات.",
    )
    if (figure := next((path for path in figure_paths if path.name == "02_top_categories.png"), None)):
        add_figure(document, figure)

    add_heading(document, "3. اختلافات وصفية بين أنواع المراكز", level=2)
    add_text(
        document,
        "يعرض الرسم متوسط إجمالي الطلب لكل مركز بحسب نوعه، ويظهر TYPE_B أعلى متوسط بين الأنواع "
        "المعروضة. هذا وصف للمراكز في البيانات المتاحة ولا يحدد سبب الفروق أو يثبت أن نوع المركز "
        "هو العامل المؤثر.",
    )
    if (figure := next((path for path in figure_paths if path.name == "03_center_type_average.png"), None)):
        add_figure(document, figure)

    add_heading(document, "4. الترويج والطلبات", level=2)
    promotion_text = analysis_cells.get("80b44d5a", "")
    email_row = re.search(
        r"emailer_for_promotion\s+(\d+)\s+[\d.]+\s+[\d.]+\s+([\d.]+)\s+([\d.]+)\s+[\d.]+\s+(\d+)",
        promotion_text,
    )
    homepage_row = re.search(
        r"homepage_featured\s+(\d+)\s+[\d.]+\s+[\d.]+\s+([\d.]+)\s+([\d.]+)\s+[\d.]+\s+(\d+)",
        promotion_text,
    )
    if email_row and homepage_row:
        add_text(
            document,
            f"في المقارنات المطابقة للمركز والوجبة، كان وسيط الطلبات مع الترويج بالبريد "
            f"{float(email_row.group(3)):,.0f} مقابل {float(email_row.group(2)):,.0f} من دونه، "
            f"ومع الظهور في الصفحة الرئيسية {float(homepage_row.group(3)):,.0f} مقابل "
            f"{float(homepage_row.group(2)):,.0f} من دونه. شملت المقارنة "
            f"{email_row.group(1)} زوجًا للبريد و{homepage_row.group(1)} زوجًا للصفحة الرئيسية. "
            "هذه ارتباطات وصفية؛ تكرار الأسابيع وعوامل أخرى محتملة يحدان من تفسيرها، ولا تثبت "
            "أن الترويج تسبب في الزيادة.",
        )
    else:
        add_text(
            document,
            "تقارن الصورة الطلبات مع الترويج ومن دونه. لم تتوفر في المخرجات المحفوظة أرقام "
            "المقارنة المطابقة، لذلك لا يورد التقرير أرقامًا لها.",
        )
    if (figure := next((path for path in figure_paths if path.name == "04_promotion_medians.png"), None)):
        add_figure(document, figure)

    add_heading(document, "5. الخصم والطلبات", level=2)
    add_text(
        document,
        f"بلغ ارتباط الرتب بين نسبة الخصم وعدد الطلبات {float(results['rho']):.3f} عبر "
        f"{results['valid_discount_rows']} مشاهدة؛ وهو ارتباط موجب ضعيف. لا يعني ذلك أن زيادة "
        "الخصم ستؤدي حتمًا إلى زيادة الطلب، لأن اختلاف الوجبات والمراكز والأسابيع قد يرتبط "
        "بالقيم المرصودة.",
    )
    if (figure := next((path for path in figure_paths if path.name == "05_discount_order_association.png"), None)):
        add_figure(document, figure)

    known_figures = {path.name for path in figure_paths if path.name in IMAGE_CAPTIONS}
    for figure in figure_paths:
        if figure.name not in known_figures:
            add_heading(document, "رسم إضافي من نتائج التحليل", level=2)
            add_figure(document, figure)

    add_heading(document, "نتائج النموذج التنبؤي")
    add_text(
        document,
        "دُرّب نموذج الانحدار الخطي على بيانات التدريب، ثم اختُبر على آخر 10 أسابيع موسومة "
        f"متاحة ({results['validation_weeks']}) قبل إعادة تدريبه على كامل بيانات التدريب. "
        f"بلغ جذر متوسط مربع الخطأ (RMSE) {results['rmse']} طلبًا لكل سجل مركز-وجبة"
        + (f"، وبلغ متوسط الخطأ المطلق (MAE) {results['mae']} طلبًا." if results["mae"] else "."),
    )
    add_summary_table(document, results)
    add_heading(document, "كيف نقرأ RMSE؟", level=2)
    add_text(
        document,
        f"يعطي RMSE مقياسًا لحجم الفرق بين الطلب الفعلي والمتوقع بوحدة الطلبات. "
        f"قيمته {results['rmse']} تعادل نحو {error_ratio:.1f}% من متوسط الطلب "
        f"({mean_target:,.2f}) في سجل التدريب. لذلك ينبغي النظر إليه كخطأ ملحوظ نسبةً إلى "
        "المتوسط، لا كدليل كافٍ وحده لقبول النموذج أو رفضه؛ فالقبول يعتمد على هامش الخطأ "
        "الذي تستطيع قرارات التشغيل تحمله، وعلى مقارنته بطريقة توقع مرجعية.",
    )
    add_text(
        document,
        f"مثال توضيحي افتراضي: إذا كان الطلب الفعلي لسجل وجبة-مركز 500 طلب، فإن فرقًا يقارب "
        f"{rmse_value:.0f} طلبًا يعني توقعًا يقارب {500 - rmse_value:.0f} أو "
        f"{500 + rmse_value:.0f} طلبًا. هذا المثال يشرح حجم المقياس ولا يمثل نتيجة فعلية. "
        "كما أن RMSE يرفع وزن الأخطاء الكبيرة، لذلك ليس متوسطًا حسابيًا بسيطًا لكل الأخطاء.",
    )

    add_heading(document, "التوصيات")
    add_bullet(document, "ابدأوا بتجربة محدودة للتوقعات على مستوى المركز والوجبة، وقارنوا التوقع بالطلب الفعلي قبل تغيير كميات التوريد.")
    add_bullet(document, "أعطوا الفئات الأعلى طلبًا أولوية في مراجعة خطط التوفر، مع التحقق من بيانات المخزون والهدر قبل اتخاذ قرارات شراء.")
    add_bullet(document, "قارنوا النموذج بطريقة مرجعية بسيطة وحددوا مسبقًا هامش الخطأ المقبول لكل قرار تشغيلي.")
    add_bullet(document, "اختبروا أثر العروض والأسعار بتجارب أو تصميم مضبوط قبل التعامل مع الارتباطات المرصودة كتوقع لزيادة الطلب.")

    add_heading(document, "ملاحظات وحدود التحليل")
    add_bullet(document, "البيانات تقيس الطلبات المسجلة، ولا تتضمن قياسًا مباشرًا للمخزون أو الهدر أو نفاد المخزون.")
    add_bullet(document, "أرقام الأسابيع ليست تواريخ تقويمية؛ لا يمكن نسبة التغيرات إلى موسم أو عطلة من هذه البيانات وحدها.")
    add_bullet(document, "مقارنات الترويج والخصم رصدية ولا تثبت السببية، كما أن المشاهدات الأسبوعية المتكررة تجعل الاختبارات الاستكشافية بحاجة إلى حذر.")
    add_bullet(document, "قيمة RMSE محسوبة على أسابيع تحقق من التدريب؛ لا تتوفر طلبات فعلية لأسابيع الاختبار لحساب خطأ التنبؤ عليها.")
    add_bullet(document, "التقرير يعرض نتائج مجمعة ولا يحاول تحديد هوية الشركة أو أي عميل أو فرد.")

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.clear()
        set_paragraph_rtl(footer, WD_ALIGN_PARAGRAPH.CENTER)
        add_page_number(footer)
    set_document_rtl(document)

    output_path = project_root / "reports" / "client_report_U.docx"
    document.save(output_path)
    verify_report(output_path, len(figure_paths), str(results["rmse"]))
    return output_path


def verify_report(path: Path, expected_images: int, rmse_text: str) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"لم يتم إنشاء ملف التقرير: {path}")
    document = Document(path)
    if len(document.inline_shapes) != expected_images:
        raise RuntimeError(
            f"عدد الصور داخل التقرير ({len(document.inline_shapes)}) لا يطابق "
            f"عدد الصور في المجلد ({expected_images})."
        )
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for table in document.tables:
        full_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    if rmse_text not in full_text or "RMSE" not in full_text or "مثال توضيحي افتراضي" not in full_text:
        raise RuntimeError("التحقق من قيمة RMSE أو شرحها في التقرير لم ينجح.")
    paragraphs = list(document.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        bidi = p_pr.find(qn("w:bidi"))
        if bidi is None or bidi.get(qn("w:val")) != "1":
            raise RuntimeError("توجد فقرة في المستند أو تذييله لا تحمل اتجاه RTL.")
    for table in document.tables:
        bidi = table._tbl.tblPr.find(qn("w:bidiVisual"))
        if bidi is None or bidi.get(qn("w:val")) != "1":
            raise RuntimeError("يوجد جدول في المستند لا يحمل اتجاه RTL.")
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    bidi = paragraph._p.get_or_add_pPr().find(qn("w:bidi"))
                    if bidi is None or bidi.get(qn("w:val")) != "1":
                        raise RuntimeError("توجد فقرة جدول لا تحمل اتجاه RTL.")


def main() -> None:
    parser = argparse.ArgumentParser(description="إنشاء تقرير العميل العربي بصيغة Word.")
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[3],
        help="مسار جذر المشروع (اختياري).",
    )
    arguments = parser.parse_args()
    project_root = arguments.project_root.resolve()
    output_path = make_report(project_root)
    print(f"تم إنشاء التقرير والتحقق منه: {output_path}")


if __name__ == "__main__":
    main()
